"""Tests for the centralized Excel presentation system."""

from datetime import datetime
from io import BytesIO

import openpyxl
import pandas as pd
from xlsxwriter import Workbook

from retailflow.reporting.formatting import (
    DEFAULT_THEME,
    configure_print_layout,
    create_report_formats,
    write_dataframe_table,
    write_internal_link,
)


def test_theme_has_restrained_management_palette() -> None:
    assert DEFAULT_THEME.primary == "#17365D"
    assert DEFAULT_THEME.white == "#FFFFFF"
    assert DEFAULT_THEME.positive != DEFAULT_THEME.warning != DEFAULT_THEME.error
    assert DEFAULT_THEME.font_name == "Aptos"


def test_table_formats_navigation_and_print_settings() -> None:
    stream = BytesIO()
    workbook = Workbook(stream, {"in_memory": True})
    formats = create_report_formats(workbook, "EUR")
    worksheet = workbook.add_worksheet("Detail")
    frame = pd.DataFrame(
        {
            "order_date": [pd.Timestamp("2025-01-01")],
            "net_revenue": [125.5],
            "gross_margin_percent": [33.3],
        }
    )
    write_internal_link(worksheet, "A1", "Summary", "Back", formats)
    write_dataframe_table(worksheet, frame, 2, 0, formats, "FormattedData", freeze_at=(3, 0))
    configure_print_layout(
        worksheet,
        report_id="TEST-1",
        generated_at=datetime(2025, 1, 2),
        last_row=5,
        last_column=2,
        repeat_header_row=2,
    )
    workbook.add_worksheet("Summary")
    workbook.close()

    stream.seek(0)
    loaded = openpyxl.load_workbook(stream)
    sheet = loaded["Detail"]
    assert sheet["A1"].hyperlink is not None
    assert sheet.freeze_panes == "A4"
    assert "EUR" in sheet["B4"].number_format
    assert sheet["C4"].number_format == '0.0"%"'
    assert sheet.tables
    assert sheet.print_area
    assert sheet.print_title_rows == "$3:$3"
